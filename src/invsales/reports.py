"""Manager-facing deliverables: a weekly action pack as an Excel workbook (built in memory)."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HDR = PatternFill("solid", fgColor="1F3A5F")


def _sheet(wb: Workbook, title: str, df: pd.DataFrame, note: str | None = None, money: tuple[str, ...] = ()) -> None:
    ws = wb.create_sheet(title[:31])
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(italic=True, color="555555")
        r0 = 3
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=r0, column=j, value=str(col).replace("_", " ").title())
        c.fill, c.font, c.alignment = HDR, Font(bold=True, color="FFFFFF"), Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(34, len(str(col)) + 4))
    for i, row in enumerate(df.itertuples(index=False), start=r0 + 1):
        for j, v in enumerate(row, start=1):
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            elif v is not None and pd.isna(v):
                v = None
            c = ws.cell(row=i, column=j, value=v)
            if df.columns[j - 1] in money:
                c.number_format = "#,##0.00"
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    if len(df):
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(df.columns))}{r0 + len(df)}"


def action_pack(*, repl: pd.DataFrame, health: dict, abc: pd.DataFrame, transfers: pd.DataFrame,
                assumptions: list[str], title: str = "Weekly inventory action pack",
                risk_threshold: float = 50.0) -> bytes:
    """Excel workbook a purchasing/warehouse manager can act on: what to order, what may stock out, what to move,
    what is overstocked, plus the assumptions the numbers rest on."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"], ws["A1"].font = title, Font(bold=True, size=14)
    rows = [("Stock value (at cost)", health["stock_value"]), ("Excess stock value (above order-up-to level)", health["excess_value"]),
            ("Excess as % of stock value", health["excess_pct"]), ("Dead stock items (no sales in 90 days)", health["dead_stock_items"]),
            ("Dead stock value", health["dead_stock_value"]),
            (f"Items with stockout risk >= {risk_threshold:.0f}% during lead time", health["at_risk_items"]),
            ("Median days of cover", health["median_days_of_cover"]), ("Suggested purchase order value (at cost)", health["suggested_po_value"])]
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
    r = 3 + len(rows) + 1
    ws.cell(row=r, column=1, value="Assumptions and limits").font = Font(bold=True)
    for k, a in enumerate(assumptions, start=r + 1):
        ws.cell(row=k, column=1, value=a)
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 18

    due = repl[repl.below_rop].sort_values("suggested_order_cost", ascending=False)
    _sheet(wb, "Orders to place", due[["warehouse_name", "product_id", "product_name", "current_stock", "reorder_point", "eoq",
                                       "suggested_order_qty", "suggested_order_cost", "stockout_risk_pct", "lead_time_days"]],
           "Items at/below the recommended reorder point. Order qty = reorder point + EOQ - stock.", money=("suggested_order_cost",))
    risk = repl[repl.stockout_risk_pct >= risk_threshold].sort_values("stockout_risk_pct", ascending=False)
    _sheet(wb, "Stockout risk", risk[["warehouse_name", "product_id", "product_name", "current_stock", "mean_daily", "lead_time_days",
                                      "stockout_risk_pct", "days_of_cover"]],
           "Probability that demand during the supplier lead time exceeds stock on hand (normal approximation).")
    _sheet(wb, "Transfers", transfers, "Move surplus stock instead of buying: surplus is stock above reorder point + EOQ.",
           money=("value_at_cost",))
    excess = repl[repl.excess_value > 0].sort_values("excess_value", ascending=False)
    _sheet(wb, "Excess stock", excess[["warehouse_name", "product_id", "product_name", "current_stock", "order_up_to", "excess_units",
                                       "excess_value", "days_of_cover"]],
           "Stock above the order-up-to level: cash tied up beyond what the policy needs.", money=("excess_value",))
    _sheet(wb, "ABC-XYZ", abc[["product_id", "product_name", "category", "revenue", "revenue_share_pct", "abc", "demand_cv", "xyz",
                               "segment"]], "A/B/C = value class, X/Y/Z = demand variability.", money=("revenue",))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
