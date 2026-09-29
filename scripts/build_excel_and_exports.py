"""Builds (1) powerbi/*.csv model tables and (2) excel/Inventory_Sales_Summary.xlsx from data/demo/clean."""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from invsales import analytics
from invsales.config import get_settings
from invsales.repository import load_clean

ROOT = Path(__file__).resolve().parent.parent
S = get_settings()
t = load_clean(ROOT / "data" / "demo" / "clean")
sales, products, inventory, restocks, warehouses = (t[k] for k in ("sales", "products", "inventory", "restocks", "warehouses"))

# ---- Power BI model tables (names match powerbi/Inventory.SemanticModel) ---------------------------
pb = ROOT / "powerbi"
sales.to_csv(pb / "fact_sales.csv", index=False, date_format="%Y-%m-%d")
inventory.to_csv(pb / "fact_inventory.csv", index=False)
restocks.to_csv(pb / "fact_restocks.csv", index=False, date_format="%Y-%m-%d")
products.to_csv(pb / "dim_products.csv", index=False)
warehouses.to_csv(pb / "dim_warehouses.csv", index=False)

# ---- Excel workbook -----------------------------------------------------------------------------------
wname = warehouses.set_index("warehouse_id").warehouse_name
df = sales.merge(products[["product_id", "product_name", "category"]], on="product_id")
df["warehouse"] = df.warehouse_id.map(wname)
df["year"] = df.order_date.dt.year
df["month"] = df.order_date.dt.to_period("M").dt.to_timestamp()      # real dates: SUMIFS mis-coerces "2024-01" text criteria
cols = ["order_id", "order_date", "year", "month", "warehouse", "product_id", "product_name", "category", "quantity",
        "unit_price", "revenue"]
df = df[cols]
LET = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}
N = len(df) + 1

HDR, WHITE = PatternFill("solid", fgColor="1F3A5F"), Font(bold=True, color="FFFFFF")


def header(ws, row=1):
    for c in ws[row]:
        c.fill, c.font, c.alignment = HDR, WHITE, Alignment(horizontal="center")


def widths(ws, w=16):
    for i in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def rng(col):
    return f"Sales_Data!${LET[col]}$2:${LET[col]}${N}"


def fmt(ws, cols, number_format, min_row=2):
    for row in ws.iter_rows(min_row=min_row):
        for c in row:
            if c.column_letter in cols:
                c.number_format = number_format


wb = Workbook()
wd = wb.active
wd.title = "Sales_Data"
wd.append(cols)
for r in df.itertuples(index=False):
    wd.append([r.order_id, r.order_date.to_pydatetime(), r.year, r.month.to_pydatetime(), r.warehouse, r.product_id,
               r.product_name, r.category, r.quantity, r.unit_price, r.revenue])
header(wd)
widths(wd, 15)
wd.freeze_panes = "A2"
fmt(wd, {"B"}, "yyyy-mm-dd")
fmt(wd, {"D"}, "mmm yyyy")
tbl = Table(displayName="SalesData", ref=f"A1:{get_column_letter(len(cols))}{N}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
wd.add_table(tbl)

# Category x year
ws = wb.create_sheet("Category_x_Year")
years = sorted(df.year.unique())
cats = sorted(df.category.unique())
ws.append(["Category", *years, "Total", "Share %"])
last_year_col = get_column_letter(1 + len(years))
tot_col = get_column_letter(2 + len(years))
for i, c in enumerate(cats, start=2):
    row = [c] + [f"=SUMIFS({rng('revenue')},{rng('category')},$A{i},{rng('year')},{get_column_letter(2 + j)}$1)" for j in range(len(years))]
    ws.append([*row, f"=SUM(B{i}:{last_year_col}{i})", f"={tot_col}{i}/${tot_col}${len(cats) + 2}"])
tr = len(cats) + 2
ws.append(["Grand Total"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{tr - 1})" for c in range(2, 3 + len(years))]
          + [f"={tot_col}{tr}/${tot_col}${tr}"])
header(ws)
widths(ws, 18)
fmt(ws, {get_column_letter(c) for c in range(2, 3 + len(years))}, "#,##0")
fmt(ws, {get_column_letter(3 + len(years))}, "0.0%")
for c in ws[tr]:
    c.font = Font(bold=True)

# Monthly sales + chart
wm = wb.create_sheet("Monthly_Sales")
wm.append(["Month", "Revenue", "Units Sold"])
for i, m in enumerate(sorted(df.month.unique()), start=2):
    wm.append([m.to_pydatetime(), f"=SUMIFS({rng('revenue')},{rng('month')},$A{i})", f"=SUMIFS({rng('quantity')},{rng('month')},$A{i})"])
header(wm)
widths(wm)
fmt(wm, {"A"}, "mmm yyyy")
fmt(wm, {"B", "C"}, "#,##0")
ch = LineChart()
ch.title, ch.height, ch.width = "Monthly Sales", 8, 18
ch.add_data(Reference(wm, min_col=2, min_row=1, max_row=wm.max_row), titles_from_data=True)
ch.set_categories(Reference(wm, min_col=1, min_row=2, max_row=wm.max_row))
wm.add_chart(ch, "E2")

# Warehouse summary
ww = wb.create_sheet("Warehouse_Summary")
ww.append(["Warehouse", "Region", "Revenue", "Units Sold", "SKUs at/below reorder level"])
for i, w in enumerate(warehouses.itertuples(), start=2):
    ww.append([w.warehouse_name, w.region, f"=SUMIFS({rng('revenue')},{rng('warehouse')},$A{i})",
               f"=SUMIFS({rng('quantity')},{rng('warehouse')},$A{i})", f'=COUNTIFS(Inventory!$A:$A,$A{i},Inventory!$F:$F,"REORDER")'])
header(ww)
widths(ww, 22)
fmt(ww, {"C", "D"}, "#,##0")
bc = BarChart()
bc.title, bc.height, bc.width = "Revenue by warehouse", 8, 14
bc.add_data(Reference(ww, min_col=3, min_row=1, max_row=ww.max_row), titles_from_data=True)
bc.set_categories(Reference(ww, min_col=1, min_row=2, max_row=ww.max_row))
ww.add_chart(bc, "G2")

# Inventory snapshot with reorder flag
wi = wb.create_sheet("Inventory")
wi.append(["Warehouse", "Product ID", "Product", "Current Stock", "Reorder Level", "Status"])
inv = inventory.merge(products[["product_id", "product_name"]], on="product_id").sort_values(["warehouse_id", "product_id"])
for i, r in enumerate(inv.itertuples(), start=2):
    wi.append([wname[r.warehouse_id], r.product_id, r.product_name, r.current_stock, r.reorder_level, f'=IF(D{i}<=E{i},"REORDER","OK")'])
header(wi)
widths(wi, 20)
wi.freeze_panes = "A2"
wi.auto_filter.ref = f"A1:F{wi.max_row}"
wi.conditional_formatting.add(f"F2:F{wi.max_row}", CellIsRule(operator="equal", formula=['"REORDER"'],
                              fill=PatternFill("solid", bgColor="F8CBAD"), font=Font(bold=True, color="9C0006")))

# Replenishment recommendations (computed in Python - see src/invsales/analytics.py)
repl = analytics.replenishment(sales, inventory, products, warehouses, S.service_level_z, S.ordering_cost, S.holding_rate)
wr = wb.create_sheet("Replenishment (Python)")
rcols = ["warehouse_name", "product_name", "current_stock", "reorder_level", "reorder_point", "safety_stock", "eoq",
         "suggested_order_qty", "suggested_order_cost"]
wr.append(["Warehouse", "Product", "Stock", "Current reorder level", "Recommended ROP", "Safety stock", "EOQ", "Order qty", "Order cost"])
for r in repl[repl.below_rop].sort_values("suggested_order_cost", ascending=False)[rcols].itertuples(index=False):
    wr.append(list(r))
header(wr)
widths(wr, 18)
wr.cell(row=wr.max_row + 2, column=1, value="Static values from the Python model; only items at/below the recommended reorder point.")

order = ["Category_x_Year", "Monthly_Sales", "Warehouse_Summary", "Inventory", "Replenishment (Python)", "Sales_Data"]
wb._sheets = [wb[n] for n in order]
wb.save(ROOT / "excel" / "Inventory_Sales_Summary.xlsx")
print("done: powerbi CSVs + excel")
