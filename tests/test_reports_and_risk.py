import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from invsales import analytics, kpis
from invsales.cli import main
from invsales.reports import action_pack
from invsales.validate import validate


def test_stockout_risk_is_monotonic_in_stock_and_bounded(sales, inventory, products, warehouses):
    lo = analytics.replenishment(sales, inventory.assign(current_stock=0), products, warehouses)
    hi = analytics.replenishment(sales, inventory.assign(current_stock=10_000), products, warehouses)
    assert lo.stockout_risk_pct.between(0, 100).all() and hi.stockout_risk_pct.between(0, 100).all()
    assert (lo.stockout_risk_pct >= hi.stockout_risk_pct).all() and hi.stockout_risk_pct.max() == 0
    demand = lo[lo.mean_daily > 0]
    assert (demand.stockout_risk_pct == 100).all()               # no stock and positive demand: certain shortage


def test_what_if_delay_raises_reorder_point_and_risk(sales, inventory, products, warehouses):
    base = analytics.replenishment(sales, inventory, products, warehouses)
    slow = analytics.replenishment(sales, inventory, products, warehouses, lead_time_extra_days=10)
    assert (slow.reorder_point >= base.reorder_point).all() and (slow.reorder_point > base.reorder_point).any()
    assert (slow.stockout_risk_pct >= base.stockout_risk_pct - 1e-9).all()
    assert slow.suggested_order_cost.sum() >= base.suggested_order_cost.sum()


def test_higher_service_level_means_more_safety_stock(sales, inventory, products, warehouses):
    a = analytics.replenishment(sales, inventory, products, warehouses, z=1.28)
    b = analytics.replenishment(sales, inventory, products, warehouses, z=2.33)
    assert (b.safety_stock >= a.safety_stock).all() and b.safety_stock.sum() > a.safety_stock.sum()


def test_excess_and_dead_stock(sales, inventory, products, warehouses):
    inv = inventory.copy()
    inv.loc[(inv.warehouse_id == "W2") & (inv.product_id == "C"), "current_stock"] = 500      # stocked, never sold in W2
    repl = analytics.replenishment(sales, inv, products, warehouses)
    row = repl[(repl.warehouse_id == "W2") & (repl.product_id == "C")].iloc[0]
    assert row.excess_units == 500 and row.excess_value == 500 * 3.0
    h = analytics.inventory_health(sales, repl)
    assert h["dead_stock_items"] >= 1 and h["dead_stock_value"] >= 1500
    assert h["stock_value"] == pytest.approx(repl.stock_value.sum()) and 0 <= h["excess_pct"] <= 100


def test_turnover_without_receipts_history_uses_current_stock(sales, inventory, products):
    empty = pd.DataFrame({"restock_date": pd.to_datetime([]), "warehouse_id": [], "product_id": [], "quantity": []})
    k = kpis.stock_turnover(sales, inventory, empty, products).set_index(["warehouse_id", "product_id"])
    assert k.loc[("W1", "A"), "stock_turnover"] == round(30 / 5, 2)                              # avg stock == current stock


def test_validate_non_strict_allows_missing_stock(tables):
    lean = {**tables, "inventory": tables["inventory"].iloc[0:0], "restocks": tables["restocks"].iloc[0:0]}
    assert validate(lean, strict=False) == [] and any("empty table" in p for p in validate(lean))


def test_action_pack_workbook(sales, inventory, products, warehouses):
    repl = analytics.replenishment(sales, inventory, products, warehouses)
    pack = action_pack(repl=repl, health=analytics.inventory_health(sales, repl), abc=analytics.abc_xyz(sales, products),
                       transfers=analytics.transfer_suggestions(repl, products), assumptions=["test assumption"])
    wb = load_workbook(io.BytesIO(pack))
    assert wb.sheetnames == ["Summary", "Orders to place", "Stockout risk", "Transfers", "Excess stock", "ABC-XYZ"]
    assert wb["Summary"]["A1"].value == "Weekly inventory action pack"
    texts = [c.value for row in wb["Summary"].iter_rows() for c in row if c.value]
    assert "test assumption" in texts and any("Suggested purchase order value" in str(t) for t in texts)
    orders = wb["Orders to place"]
    assert orders.max_row == 3 + int(repl.below_rop.sum())                                       # note row + blank + header + rows


def test_cli_import_and_templates_end_to_end(tmp_path, capsys):
    assert main(["templates"]) == 0
    from invsales.templates import write_templates
    files = write_templates(tmp_path / "t")
    out = tmp_path / "out"
    rc = main(["import", "--sales", files[0], "--stock", files[1], "--products", files[2], "--out", str(out)])
    assert rc == 0 and (out / "sales.parquet").exists() and (out / "inventory.parquet").exists()
    said = capsys.readouterr().out
    assert "column mapping" in said and "days of history" in said
    inv = pd.read_parquet(out / "inventory.parquet")
    assert len(inv) == 4 and inv.reorder_level.dtype.kind == "i"


def test_cli_import_fails_cleanly_on_unusable_file(tmp_path, capsys):
    bad = tmp_path / "bad.csv"
    bad.write_text("foo,bar\n1,2\n")
    assert main(["import", "--sales", str(bad), "--out", str(tmp_path / "o")]) == 1
    assert "ERROR" in capsys.readouterr().out and not (tmp_path / "o").exists()
