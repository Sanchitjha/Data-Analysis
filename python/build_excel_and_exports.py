"""Builds (1) excel/Inventory_Sales_Summary.xlsx and (2) powerbi/*.csv model tables from data/clean."""
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

ROOT = Path(__file__).resolve().parent.parent
sales = pd.read_csv(ROOT/"data/clean/sales.csv", parse_dates=["order_date"])
prod = pd.read_csv(ROOT/"data/clean/products.csv")
rest = pd.read_csv(ROOT/"data/clean/restocks.csv", parse_dates=["restock_date"])

# ---- Power BI model tables --------------------------------------------------
pb = ROOT/"powerbi"
sales.to_csv(pb/"fact_sales.csv", index=False, date_format="%Y-%m-%d")
prod.to_csv(pb/"dim_products.csv", index=False)
rest.to_csv(pb/"fact_restocks.csv", index=False, date_format="%Y-%m-%d")

# ---- Excel workbook ---------------------------------------------------------
df = sales.merge(prod[["product_id", "product_name", "category"]], on="product_id")
df["year"] = df.order_date.dt.year
df["month"] = df.order_date.dt.to_period("M").dt.to_timestamp()   # real dates: SUMIFS mis-coerces "2024-01" text criteria
df = df[["order_id", "order_date", "year", "month", "product_id", "product_name", "category", "store", "quantity", "unit_price", "revenue"]]

wb = Workbook()
HDR = PatternFill("solid", fgColor="1F3A5F"); WHITE = Font(bold=True, color="FFFFFF")
def header(ws, row=1):
    for c in ws[row]:
        c.fill, c.font, c.alignment = HDR, WHITE, Alignment(horizontal="center")
def widths(ws, w=16):
    for i in range(1, ws.max_column + 1): ws.column_dimensions[get_column_letter(i)].width = w

# Data sheet (an Excel Table -> ideal PivotTable source)
wd = wb.active; wd.title = "Sales_Data"
wd.append(list(df.columns))
for r in df.itertuples(index=False):
    wd.append([r.order_id, r.order_date.to_pydatetime(), r.year, r.month.to_pydatetime(), r.product_id, r.product_name, r.category, r.store, r.quantity, r.unit_price, r.revenue])
for row in wd.iter_rows(min_row=2, min_col=2, max_col=4):
    row[0].number_format = "yyyy-mm-dd"; row[2].number_format = "mmm yyyy"
header(wd); widths(wd, 15); wd.freeze_panes = "A2"
from openpyxl.worksheet.table import Table, TableStyleInfo
t = Table(displayName="SalesData", ref=f"A1:K{len(df)+1}"); t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
wd.add_table(t)
N = len(df) + 1
rng = lambda col: f"Sales_Data!${col}$2:${col}${N}"

# Summary by category x year (SUMIFS = same numbers a PivotTable would show)
ws = wb.create_sheet("Category_x_Year")
ws.append(["Category", 2024, 2025, "Total", "Share %"])
cats = sorted(df.category.unique())
for i, c in enumerate(cats, start=2):
    ws.append([c, f"=SUMIFS({rng('K')},{rng('G')},$A{i},{rng('C')},B$1)", f"=SUMIFS({rng('K')},{rng('G')},$A{i},{rng('C')},C$1)", f"=B{i}+C{i}", f"=D{i}/$D${len(cats)+2}"])
tr = len(cats) + 2
ws.append(["Grand Total"] + [f"=SUM({c}2:{c}{tr-1})" for c in "BCD"] + [f"=D{tr}/$D${tr}"])
header(ws); widths(ws, 18)
for r in ws.iter_rows(min_row=2, min_col=2, max_col=4):
    for c in r: c.number_format = "#,##0"
for r in ws.iter_rows(min_row=2, min_col=5, max_col=5):
    for c in r: c.number_format = "0.0%"
for c in ws[tr]: c.font = Font(bold=True)

# Monthly sales
wm = wb.create_sheet("Monthly_Sales")
wm.append(["Month", "Revenue", "Units Sold"])
for i, m in enumerate(sorted(df.month.unique()), start=2):
    wm.append([m.to_pydatetime(), f"=SUMIFS({rng('K')},{rng('D')},$A{i})", f"=SUMIFS({rng('I')},{rng('D')},$A{i})"])
header(wm); widths(wm, 16)
for row in wm.iter_rows(min_row=2, max_col=1): row[0].number_format = "mmm yyyy"
for r in wm.iter_rows(min_row=2, min_col=2, max_col=3):
    for c in r: c.number_format = "#,##0"
from openpyxl.chart import LineChart, Reference
ch = LineChart(); ch.title = "Monthly Sales"; ch.height, ch.width = 8, 18
ch.add_data(Reference(wm, min_col=2, min_row=1, max_row=wm.max_row), titles_from_data=True)
ch.set_categories(Reference(wm, min_col=1, min_row=2, max_row=wm.max_row)); wm.add_chart(ch, "E2")

# Product summary: units, revenue, stock, reorder flag
wp = wb.create_sheet("Product_Summary")
wp.append(["Product ID", "Product", "Category", "Units Sold", "Revenue", "Current Stock", "Reorder Level", "Reorder Alert"])
for i, p in enumerate(prod.sort_values("product_id").itertuples(), start=2):
    wp.append([p.product_id, p.product_name, p.category,
               f"=SUMIFS({rng('I')},{rng('E')},$A{i})", f"=SUMIFS({rng('K')},{rng('E')},$A{i})",
               p.current_stock, p.reorder_level, f'=IF(F{i}<=G{i},"REORDER","OK")'])
header(wp); widths(wp, 18); wp.freeze_panes = "A2"
wp.conditional_formatting.add(f"H2:H{wp.max_row}", CellIsRule(operator="equal", formula=['"REORDER"'], fill=PatternFill("solid", bgColor="F8CBAD"), font=Font(bold=True, color="9C0006")))
wp.auto_filter.ref = f"A1:H{wp.max_row}"

wb.move_sheet("Category_x_Year", offset=-1); wb.move_sheet("Monthly_Sales", offset=-1); wb.move_sheet("Product_Summary", offset=-1)
wb.save(ROOT/"excel/Inventory_Sales_Summary.xlsx")
print("done")
